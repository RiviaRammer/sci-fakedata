"""CSV/TSV, Excel and NumPy input; no network access or pandas dependency."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from io import StringIO
from pathlib import Path
from typing import Any

import numpy as np


@dataclass
class Table:
    values: np.ndarray
    text: np.ndarray | None = None
    source: str | None = None
    notes: tuple[str, ...] = ()


def _number(cell: Any) -> float:
    if cell is None or isinstance(cell, (bool, np.bool_)):
        return float("nan")
    try:
        return float(str(cell).strip())
    except (TypeError, ValueError):
        return float("nan")


def _table(rows, *, header="auto", skiprows=0, columns=None, source=None):
    if header not in ("auto", True, False):
        raise ValueError("header must be 'auto', True or False")
    if not isinstance(skiprows, int) or skiprows < 0:
        raise ValueError("skiprows must be a non-negative integer")
    rows = [list(row) for row in rows][skiprows:]
    if not rows:
        raise ValueError("input has no rows")
    width = max(map(len, rows))
    if not width:
        raise ValueError("input has no columns")
    # Pad ragged tables instead of shifting or dropping missing positions.
    rows = [row + [None] * (width - len(row)) for row in rows]
    if columns is not None:
        indices = list(columns)
        if not indices or any(not isinstance(i, int) or not 0 <= i < width for i in indices):
            raise ValueError("columns must contain valid zero-based column indices")
        rows = [[row[i] for i in indices] for row in rows]
    text = np.array([["" if cell is None else str(cell).strip() for cell in row] for row in rows], dtype=object)
    values = np.array([[_number(cell) for cell in row] for row in rows], dtype=float)
    skip_header = header is True or (
        header == "auto" and not np.isfinite(values[0]).any()
        and len(values) > 1 and np.isfinite(values[1:]).any()
    )
    if skip_header:
        values, text = values[1:], text[1:]
    if not values.size or not np.isfinite(values).any():
        raise ValueError("input has no finite numeric values")
    notes = (
        "Text/blank cells are NaN; numeric cells stay in their original relative positions.",
        "Coordinates are zero-based array indices after skiprows/header/column selection.",
    )
    return Table(values, text, source, notes)


def _excel_rows(path, sheet):
    suffix = path.suffix.lower()
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ImportError("xlrd is a required dependency; reinstall sci-fakedata with its dependencies") from exc
        book = xlrd.open_workbook(path)
        try:
            selected = book.sheet_by_index(sheet) if isinstance(sheet, int) else book.sheet_by_name(sheet)
            return [selected.row_values(i) for i in range(selected.nrows)]
        finally:
            book.release_resources()
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("openpyxl is a required dependency; reinstall sci-fakedata with its dependencies") from exc
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        selected = book.worksheets[sheet] if isinstance(sheet, int) else book[sheet]
        return list(selected.values)
    finally:
        book.close()


def read(path, *, sheet=0, header="auto", delimiter=None, skiprows=0, columns=None, encoding="utf-8-sig") -> Table:
    """Read one table, retaining text precision for scan(path).

    Excel formulas use their stored cached values. Excel number formatting is
    not interpreted as instrument precision; uncached formulas become NaN.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".npy":
        values = np.load(path, allow_pickle=False)
        if values.dtype.kind not in "iuf":
            raise ValueError("NumPy input must contain real numeric values")
        return Table(np.asarray(values, dtype=float), source=path.name)
    if suffix in (".xlsx", ".xls", ".xlsm"):
        rows = _excel_rows(path, sheet)
        return _table(rows, header=header, skiprows=skiprows, columns=columns, source=f"{path.name}:{sheet}")
    if suffix not in (".csv", ".tsv", ".txt"):
        raise ValueError("supported inputs: .csv, .tsv, .txt, .xlsx, .xls, .xlsm, .npy")
    with path.open(newline="", encoding=encoding) as stream:
        if delimiter is None:
            if suffix == ".tsv":
                delimiter = "\t"
            else:
                sample = stream.read(8192)
                stream.seek(0)
                try:
                    delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
                except csv.Error:
                    # Sniffer rejects ragged rows. Count parsed field breaks
                    # while still respecting quoted delimiters.
                    candidates = ",;\t"
                    delimiter = max(candidates, key=lambda sep: sum(
                        max(0, len(row) - 1) for row in csv.reader(StringIO(sample), delimiter=sep)
                    ))
        rows = list(csv.reader(stream, delimiter=delimiter))
    return _table(rows, header=header, skiprows=skiprows, columns=columns, source=path.name)


def load(path, **kwargs) -> np.ndarray:
    """Load a file as a float NumPy ndarray. Use scan(path) to retain CSV text precision."""
    return read(path, **kwargs).values


def sheet_names(path) -> list[str]:
    path = Path(path)
    if path.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ImportError("xlrd is a required dependency; reinstall sci-fakedata with its dependencies") from exc
        book = xlrd.open_workbook(path)
        try:
            return book.sheet_names()
        finally:
            book.release_resources()
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("openpyxl is a required dependency; reinstall sci-fakedata with its dependencies") from exc
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def load_all(path, **kwargs) -> dict[str, np.ndarray]:
    """Read every numeric worksheet. Blank/text-only worksheets are omitted."""
    if "sheet" in kwargs:
        raise ValueError("load_all chooses every sheet; do not pass sheet")
    result = {}
    for name in sheet_names(path):
        try:
            result[name] = load(path, sheet=name, **kwargs)
        except ValueError as exc:
            if str(exc) not in ("input has no rows", "input has no columns", "input has no finite numeric values"):
                raise
    if not result:
        raise ValueError("workbook has no numeric worksheets")
    return result
